"""The Excel hand-out of a blind worksheet, and the way back.

What is guarded here is the measurement, not the spreadsheet: that a rater
cannot be handed answers, that their answers land on the case they were
given for, that Excel's silent conversions are refused rather than guessed
back, and that the instrument they rated under is the one the scorer sees.
"""

from __future__ import annotations

import copy
import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import appraisal as ap  # noqa: E402
import eval_agreement as ea  # noqa: E402
import worksheet_xlsx as wx  # noqa: E402

import openpyxl  # noqa: E402

CALIBRATION = ROOT / "eval" / "claim_prefill_calibration.json"
CATALOGUE = ROOT / "eval" / "catalog_second_rater.json"


def _load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


class _Workbook(unittest.TestCase):
    """Export into a temp dir; hand back paths and column positions."""

    source = CALIBRATION

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.original = _load(self.source)
        self.xlsx = self.tmp / "bogen.xlsx"
        wx.export_workbook(copy.deepcopy(self.original), self.xlsx)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def open(self):
        wb = openpyxl.load_workbook(self.xlsx)
        header = [c.value for c in wb[wx.SHEET_RATING][1]]
        return wb, {name: i + 1 for i, name in enumerate(header)}

    def sign(self, wb, rater: str = "Test", day: dt.date = dt.date(2026, 10, 1)) -> None:
        ins = wb[wx.SHEET_INSTRUCTIONS]
        ins[wx.PROTOCOL_CELLS["rater"]] = rater
        ins[wx.PROTOCOL_CELLS["labeled_at"]] = day

    def save_and_import(self, wb) -> dict:
        wb.save(self.xlsx)
        return wx.import_workbook(self.xlsx)

    def problems(self, wb) -> list[str]:
        wb.save(self.xlsx)
        with self.assertRaises(wx.WorkbookError) as caught:
            wx.import_workbook(self.xlsx)
        return caught.exception.problems


class RoundTripTests(_Workbook):
    def test_an_untouched_sheet_comes_back_changed_only_in_its_protocol(self) -> None:
        wb, _ = self.open()
        self.sign(wb)
        back = self.save_and_import(wb)
        expected = copy.deepcopy(self.original)
        expected["protocol"].update(rater="Test", labeled_at="2026-10-01", blind=True, notes="")
        self.assertEqual(back, expected)

    def test_answers_land_on_their_case_not_on_their_row(self) -> None:
        # Two rows swapped wholesale -- ID and answers together, as a
        # copy-paste would. Matching by position would cross the answers.
        wb, col = self.open()
        ws = wb[wx.SHEET_RATING]
        ws.protection.sheet = False
        first, second = self.original["labels"][0]["key"], self.original["labels"][1]["key"]
        for c in range(1, ws.max_column + 1):
            a, b = ws.cell(2, c).value, ws.cell(3, c).value
            ws.cell(2, c).value, ws.cell(3, c).value = b, a
        # Row 2 now holds the SECOND case.
        ws.cell(2, col["evidence_certainty"]).value = "strong"
        ws.cell(3, col["evidence_certainty"]).value = "very_low"
        self.sign(wb)
        back = {item["key"]: item for item in self.save_and_import(wb)["labels"]}
        self.assertEqual(back[second]["evidence_certainty"], "strong")
        self.assertEqual(back[first]["evidence_certainty"], "very_low")

    def test_the_imported_pass_scores_and_stays_a_calibration_round(self) -> None:
        wb, col = self.open()
        for row in range(2, len(self.original["labels"]) + 2):
            wb[wx.SHEET_RATING].cell(row, col["evidence_certainty"]).value = "low"
        self.sign(wb)
        back = self.save_and_import(wb)
        out = self.tmp / "back.json"
        out.write_text(json.dumps(back), encoding="utf-8")
        comparisons = ea.second_rater_comparisons(out)
        self.assertTrue(comparisons)
        for comparison in comparisons:
            with self.subTest(comparison.field):
                self.assertTrue(comparison.calibration)
                self.assertFalse(comparison.gate_ready())
        self.assertEqual(
            back["protocol"]["appraisal_method_at_rating"],
            self.original["protocol"]["appraisal_method_at_rating"],
        )

    def test_a_blank_cell_is_null_and_null_is_a_value(self) -> None:
        # effect_direction "null" means a null FINDING; an empty cell means
        # no answer. Conflating the two would turn every null result into
        # a skipped field.
        wb, col = self.open()
        ws = wb[wx.SHEET_RATING]
        ws.cell(2, col["effect_direction"]).value = "null"
        ws.cell(3, col["effect_direction"]).value = "   "
        self.sign(wb)
        labels = self.save_and_import(wb)["labels"]
        self.assertEqual(labels[0]["effect_direction"], "null")
        self.assertIsNone(labels[1]["effect_direction"])


class CatalogueRoundTripTests(_Workbook):
    source = CATALOGUE

    def test_the_large_sheet_survives_and_every_cell_fits_excel(self) -> None:
        # The catalogue worksheet is ~91,000 characters; one Excel cell
        # holds 32,767. openpyxl does not enforce that, Excel does -- by
        # truncating -- so the limit is checked here rather than trusted.
        wb, _ = self.open()
        meta = wb[wx.SHEET_META]
        self.assertGreater(int(meta["B3"].value), 1)
        for row in meta.iter_rows():
            for cell in row:
                self.assertLessEqual(len(str(cell.value or "")), 32_767)
        self.sign(wb)
        back = self.save_and_import(wb)
        self.assertEqual(back["labels"], self.original["labels"])
        # And the context the catalogue adds reaches the rater.
        self.assertIn("Titel der Quelle", [c.value for c in wb[wx.SHEET_RATING][1]])


class ExportGuardTests(unittest.TestCase):
    def test_a_sheet_with_any_answer_is_refused(self) -> None:
        sheet = _load(CALIBRATION)
        sheet["labels"][3]["study_design"] = "rct"
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(wx.WorkbookError) as caught:
                wx.export_workbook(sheet, Path(tmp) / "x.xlsx")
        self.assertIn(sheet["labels"][3]["key"], "\n".join(caught.exception.problems))

    def test_the_labelled_eval_set_cannot_be_handed_out(self) -> None:
        # The file with the answers in it is not a worksheet at all. The
        # likeliest accident is passing it instead of the blank sheet.
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(wx.WorkbookError):
                wx.export_workbook(
                    _load(ROOT / "eval" / "claim_prefill_labeled.json"), Path(tmp) / "x.xlsx"
                )

    def test_a_completed_pass_is_refused(self) -> None:
        sheet = _load(CALIBRATION)
        sheet["protocol"]["rater"] = "someone"
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(wx.WorkbookError):
                wx.export_workbook(sheet, Path(tmp) / "x.xlsx")


class ImportGuardTests(_Workbook):
    def test_a_date_in_an_age_cell_is_refused_not_converted_back(self) -> None:
        wb, col = self.open()
        wb[wx.SHEET_RATING].cell(2, col["age_range_explicit"]).value = dt.datetime(2026, 12, 10)
        self.sign(wb)
        problems = "\n".join(self.problems(wb))
        self.assertIn("als Datum", problems)
        self.assertIn("age_range_explicit", problems)

    def test_the_age_columns_are_stored_as_text(self) -> None:
        # The cause of the date problem above, prevented at the source.
        wb, col = self.open()
        ws = wb[wx.SHEET_RATING]
        for field in wx.AGE_FIELDS & set(col):
            for row in range(2, len(self.original["labels"]) + 2):
                with self.subTest(field, row=row):
                    self.assertEqual(ws.cell(row, col[field]).number_format, "@")

    def test_malformed_and_reversed_bands_are_refused(self) -> None:
        wb, col = self.open()
        ws = wb[wx.SHEET_RATING]
        ws.cell(2, col["age_range_explicit"]).value = "8 bis 10"
        ws.cell(3, col["age_range_explicit"]).value = "12-8"
        ws.cell(4, col["age_range_explicit"]).value = 10
        self.sign(wb)
        problems = self.problems(wb)
        self.assertEqual(len(problems), 3)

    def test_a_value_outside_the_vocabulary_is_named_with_row_and_field(self) -> None:
        wb, col = self.open()
        wb[wx.SHEET_RATING].cell(5, col["study_design"]).value = "RCT"
        self.sign(wb)
        problems = self.problems(wb)
        self.assertEqual(len(problems), 1)
        self.assertIn("Zeile 5", problems[0])
        self.assertIn("study_design", problems[0])
        self.assertIn('"RCT"', problems[0])

    def test_every_problem_is_reported_at_once(self) -> None:
        wb, col = self.open()
        wb[wx.SHEET_RATING].cell(2, col["claim_type"]).value = "kausal"
        wb[wx.SHEET_RATING].cell(3, col["claim_type"]).value = "kausal"
        # No rater, no date.
        self.assertEqual(len(self.problems(wb)), 4)

    def test_a_missing_or_doubled_case_is_refused(self) -> None:
        wb, _ = self.open()
        ws = wb[wx.SHEET_RATING]
        dropped = ws.cell(3, 2).value
        ws.cell(3, 2).value = ws.cell(2, 2).value
        self.sign(wb)
        problems = "\n".join(self.problems(wb))
        self.assertIn("doppelt", problems)
        self.assertIn(dropped, problems)

    def test_an_edited_hidden_original_fails_its_checksum(self) -> None:
        wb, _ = self.open()
        meta = wb[wx.SHEET_META]
        meta["A5"] = str(meta["A5"].value).replace('"1.3.0"', '"9.9.9"', 1)
        self.sign(wb)
        problems = self.problems(wb)
        self.assertIn("checksum", problems[0])

    def test_moved_answer_columns_are_refused(self) -> None:
        wb, col = self.open()
        ws = wb[wx.SHEET_RATING]
        a, b = col["claim_type"], col["study_design"]
        ws.cell(1, a).value, ws.cell(1, b).value = ws.cell(1, b).value, ws.cell(1, a).value
        self.sign(wb)
        self.assertIn("columns", self.problems(wb)[0])

    def test_blind_no_is_carried_through(self) -> None:
        wb, _ = self.open()
        self.sign(wb)
        wb[wx.SHEET_INSTRUCTIONS][wx.PROTOCOL_CELLS["blind"]] = "nein"
        self.assertIs(self.save_and_import(wb)["protocol"]["blind"], False)


class WorkbookShapeTests(_Workbook):
    def _dropdowns(self, wb, col) -> dict[str, list[str]]:
        ws = wb[wx.SHEET_RATING]
        lists = wb[wx.SHEET_LISTS]
        by_field: dict[str, list[str]] = {}
        for dv in ws.data_validations.dataValidation:
            if dv.type != "list":
                continue
            name = dv.formula1.lstrip("=")
            ref = wb.defined_names[name].attr_text.split("!")[1].replace("$", "")
            values = [c.value for row in lists[ref] for c in row]
            column = str(dv.sqref).split(":")[0].rstrip("0123456789")
            field = [f for f, i in col.items() if wx._col(i) == column][0]
            by_field[field] = values
        return by_field

    def test_every_dropdown_offers_exactly_the_vocabulary(self) -> None:
        wb, col = self.open()
        dropdowns = self._dropdowns(wb, col)
        expected = {
            "evidence_certainty": list(ap.CERTAINTY_VALUES),
            "claim_type": list(ap.CLAIM_TYPE_VALUES),
            "claim_supported_by_source": list(ap.CLAIM_SUPPORT_VALUES),
            "study_design": list(ap.STUDY_DESIGN_VALUES),
            "effect_direction": list(ap.EFFECT_DIRECTION_VALUES),
            "evidence_strength": ["low", "moderate", "strong"],
        }
        self.assertEqual(dropdowns, expected)
        # The legacy scale is checked against the labels that use it,
        # since the appraisal module no longer carries it.
        gold = {e["gold"].get("evidence_strength") for e in _load(ROOT / "eval" / "claim_prefill_labeled.json")["examples"]}
        self.assertLessEqual(gold - {None}, set(dropdowns["evidence_strength"]))

    def test_only_the_answer_cells_are_editable(self) -> None:
        wb, col = self.open()
        ws = wb[wx.SHEET_RATING]
        self.assertTrue(ws.protection.sheet)
        rated = set(self.original["protocol"]["rated_fields"])
        for name, index in col.items():
            with self.subTest(name):
                self.assertEqual(ws.cell(2, index).protection.locked, name not in rated)

    def test_the_case_id_is_kept_but_out_of_sight(self) -> None:
        # Import needs the ID; the rater should not read it. Some IDs name
        # the design ("-rct") or the document type ("policy-"), and two end
        # in "-null" although their direction is not_applicable.
        wb, col = self.open()
        ws = wb[wx.SHEET_RATING]
        self.assertEqual(col["Fall-ID"], 2)
        self.assertTrue(ws.column_dimensions["B"].hidden)
        self.assertEqual(ws.cell(2, 2).value, self.original["labels"][0]["key"])

    def test_the_progress_count_covers_exactly_the_answer_block(self) -> None:
        wb, col = self.open()
        rated = self.original["protocol"]["rated_fields"]
        first, last = wx._col(col[rated[0]]), wx._col(col[rated[-1]])
        rows = len(self.original["labels"]) + 1
        self.assertEqual(
            wb[wx.SHEET_INSTRUCTIONS]["C17"].value,
            f"=COUNTA('{wx.SHEET_RATING}'!{first}2:{last}{rows})",
        )

    def _xml(self, name: str) -> str:
        import zipfile

        with zipfile.ZipFile(self.xlsx) as z:
            return z.read(name).decode("utf-8")

    def test_excel_computes_the_count_on_open(self) -> None:
        # openpyxl writes the formula without a value. The file must tell
        # Excel to calculate on load, or the count reads empty until the
        # rater edits a cell. openpyxl sets this itself -- the test pins
        # that default on the file Excel reads, because the hand-out
        # depends on it and nothing here would notice if it changed.
        self.assertIn('fullCalcOnLoad="1"', self._xml("xl/workbook.xml"))

    def test_the_rater_meets_the_instructions_and_no_hidden_tab_is_selected(self) -> None:
        # Excel refuses a workbook whose selected tab is hidden. Also an
        # openpyxl behaviour, pinned on the raw XML for the same reason.
        import re

        wbx = self._xml("xl/workbook.xml")
        sheets = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*state="([^"]+)"', wbx)
        active = int((re.findall(r'activeTab="(\d+)"', wbx) or ["0"])[0])
        self.assertEqual(sheets[active], (wx.SHEET_INSTRUCTIONS, "visible"))
        for index, (name, state) in enumerate(sheets, start=1):
            if state != "visible":
                with self.subTest(name):
                    sheet = self._xml(f"xl/worksheets/sheet{index}.xml")
                    self.assertNotIn('tabSelected="1"', sheet)

    def test_the_rubric_sheet_carries_every_definition_verbatim(self) -> None:
        wb, _ = self.open()
        texts = {str(c.value) for row in wb[wx.SHEET_RUBRIC].iter_rows() for c in row if c.value}
        for key, value in self.original.items():
            if not key.startswith("rubrik_"):
                continue
            for text in (value.values() if isinstance(value, dict) else [value]):
                with self.subTest(key):
                    self.assertIn(text, texts)


if __name__ == "__main__":
    unittest.main()
